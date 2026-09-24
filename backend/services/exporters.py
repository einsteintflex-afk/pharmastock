# ============================================================
# REPORT EXPORT: CSV, EXCEL, PDF
# ============================================================

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .reports import Report

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value: str) -> str:
    """Neutralise spreadsheet formula injection (OWASP CSV injection)."""
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _cell(value, kind: str):
    if value is None:
        return ""
    if kind == "money":
        return round(float(value), 2)
    if kind == "number":
        return float(value)
    if kind == "int":
        return int(value)
    if kind == "datetime" and isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return _safe_text(str(value))


def _display(value, kind: str) -> str:
    cell = _cell(value, kind)
    if cell == "":
        return "—"
    if kind == "money":
        return f"{cell:,.2f}"
    if kind == "int":
        return f"{cell:,}"
    if kind == "number":
        return f"{cell:,.2f}".rstrip("0").rstrip(".")
    return str(cell)


def to_json(report: Report) -> dict:
    return {
        "report": report.key,
        "title": report.title,
        "subtitle": report.subtitle,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "columns": [{"key": k, "label": label, "type": t} for k, label, t in report.columns],
        "rows": [{k: row.get(k) for k, _, _ in report.columns} for row in report.rows],
        "summary": [{"label": label, "value": value} for label, value in report.summary],
    }


def to_csv(report: Report) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _, label, _ in report.columns])
    for row in report.rows:
        writer.writerow([_cell(row.get(k), t) for k, _, t in report.columns])
    if report.summary:
        writer.writerow([])
        for label, value in report.summary:
            writer.writerow([_safe_text(label), _safe_text(value)])
    # UTF-8 with BOM so Excel shows the currency symbol correctly.
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(report: Report, generated_by: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31]

    sheet.append([report.title])
    sheet["A1"].font = Font(size=14, bold=True)
    sheet.append([f"{report.subtitle} — generated {datetime.now():%Y-%m-%d %H:%M} by {generated_by}"])
    sheet.append([])

    header_row = sheet.max_row + 1
    sheet.append([label for _, label, _ in report.columns])
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(vertical="center")

    for row in report.rows:
        sheet.append([_cell(row.get(k), t) for k, _, t in report.columns])
        for index, (_, _, kind) in enumerate(report.columns, start=1):
            if kind == "money":
                sheet.cell(row=sheet.max_row, column=index).number_format = "#,##0.00"

    if report.summary:
        sheet.append([])
        for label, value in report.summary:
            sheet.append([_safe_text(label), _safe_text(value)])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

    for index, (key, label, kind) in enumerate(report.columns, start=1):
        width = max([len(label)] + [len(_display(r.get(key), kind)) for r in report.rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 45)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_pdf(report: Report, generated_by: str, brand: dict | None = None) -> bytes:
    """brand: the organization's branding (services/branding.profile): name,
    logo, details and footer in the header / footer of every report."""
    buffer = io.BytesIO()
    brand = brand or {}
    document = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=16 * mm, title=report.title, author=brand.get("name") or "PharmaStock",
    )
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    header_style = styles["BodyText"].clone("header", fontSize=7.5, leading=9, textColor=colors.white,
                                            fontName="Helvetica-Bold")

    def para(text, style):
        safe = str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    company = brand.get("name") or "PharmaStock"
    heading = [para(company, styles["Heading2"]),
               para(" · ".join(x for x in (brand.get("address"), brand.get("phone"), brand.get("email")) if x),
                    small)]
    if brand.get("logo_png"):
        from PIL import Image as PILImage
        from reportlab.platypus import Image
        with PILImage.open(io.BytesIO(brand["logo_png"])) as probe:
            ratio = probe.height / probe.width
        logo = Image(io.BytesIO(brand["logo_png"]), width=28 * mm, height=28 * mm * ratio)
        head = Table([[logo, heading]], colWidths=[32 * mm, None])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story = [head]
    else:
        story = heading
    story += [
        Paragraph(report.title, styles["Title"]),
        Paragraph(f"{report.subtitle}<br/>Generated {datetime.now():%Y-%m-%d %H:%M} by {generated_by}",
                  styles["BodyText"]),
        Spacer(1, 5 * mm),
    ]

    data = [[para(label, header_style) for _, label, _ in report.columns]]
    for row in report.rows:
        data.append([para(_display(row.get(k), t), small) for k, _, t in report.columns])
    if len(data) == 1:
        data.append([para("No records", small)] + [""] * (len(report.columns) - 1))

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0B7C3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)

    if report.summary:
        story.append(Spacer(1, 5 * mm))
        summary = Table([[para(label, small), para(value, small)] for label, value in report.summary],
                        colWidths=[90 * mm, 70 * mm])
        summary.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0B7C3"))]))
        story.append(summary)

    footer_text = " — ".join(x for x in (brand.get("report_footer"), brand.get("powered_by")) if x)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6B7686"))
        canvas.drawString(12 * mm, 8 * mm, f"{company} · {report.title}" + (f" · {footer_text}" if footer_text else ""))
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
